"""Validate sports.xlsx and invoke the existing direct sport generator."""

import argparse
from dataclasses import dataclass
from enum import Enum
from pathlib import Path
import re
import subprocess
import sys
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

import yaml
from openpyxl import load_workbook


VERSION = "0.1.0"
COMMON_RELATIVE = Path("framework/engine/sportEngine/common")
XLSX_NAME = "sports.xlsx"
GENERATOR_NAME = "sport_gen.py"
RULES_PATH = Path(__file__).resolve().parents[2] / "rules" / "sports_xlsx_rules.yaml"
ACTION_ALIASES = {
    "-c": "check",
    "-check": "check",
    "check": "check",
    "-g": "gen",
    "-gen": "gen",
    "gen": "gen",
}


class SportConfigError(RuntimeError):
    pass


class CheckResult(Enum):
    SKIPPED = "skipped"
    PASSED = "passed"


@dataclass(frozen=True)
class Finding:
    severity: str
    rule_id: str
    message: str


def normalize_short_options(argv: Sequence[str]) -> List[str]:
    normalized = []
    for argument in argv:
        if len(argument) >= 2 and argument[0] == "-" and argument[1] != "-" and argument[1].isalpha():
            normalized.append("-" + argument[1].lower() + argument[2:])
        else:
            normalized.append(argument)
    return normalized


def normalize_action_aliases(argv: Sequence[str]) -> List[str]:
    normalized = []
    for argument in normalize_short_options(argv):
        normalized.append(ACTION_ALIASES.get(argument.lower(), argument))
    return normalized


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Check sports.xlsx and run its existing direct generator.",
        usage="sport-config [-h] [-v] ACTION [-r PATH]",
        epilog="ACTION: -c/-C/-check/check for validation; -g/-G/-gen/gen for direct generation.",
    )
    parser.add_argument("-v", "--version", action="version", version="sport-config " + VERSION)
    parser.add_argument("command", choices=("check", "gen"), metavar="ACTION", help="-c/-C/-check/check or -g/-G/-gen/gen")
    parser.add_argument("-r", "--repo", metavar="PATH", help="repo root or a directory inside it")
    return parser


def find_repo(start: Path) -> Path:
    current = start.resolve()
    if current.is_file():
        current = current.parent
    for candidate in (current, *current.parents):
        if (candidate / ".repo").is_dir():
            return candidate
    raise SportConfigError("cannot find a repo root containing .repo from: {}".format(start))


def require_layout(repo: Path) -> Tuple[Path, Path]:
    common_dir = repo / COMMON_RELATIVE
    xlsx = common_dir / XLSX_NAME
    generator = common_dir / GENERATOR_NAME
    if not xlsx.is_file():
        raise SportConfigError("missing sports.xlsx: {}".format(xlsx))
    if not generator.is_file():
        raise SportConfigError("missing sport generator: {}".format(generator))
    if not RULES_PATH.is_file():
        raise SportConfigError("missing rule configuration: {}".format(RULES_PATH))
    return xlsx, generator


def xlsx_has_changes(repo: Path, xlsx: Path) -> bool:
    git_root = git_root_for_path(xlsx.parent)
    relative = xlsx.relative_to(git_root)
    result = git_command(git_root, ["status", "--porcelain=v1", "--untracked-files=all", "--", str(relative)])
    return bool(result.stdout.strip())


def git_root_for_path(start: Path) -> Path:
    result = git_command(start, ["rev-parse", "--show-toplevel"])
    return Path(result.stdout.strip()).resolve()


def git_command(directory: Path, arguments: Sequence[str]) -> subprocess.CompletedProcess:
    try:
        result = subprocess.run(
            ["git", "-C", str(directory), *arguments],
            check=False,
            text=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )
    except OSError as exc:
        raise SportConfigError("cannot run git: {}".format(exc))
    if result.returncode:
        raise SportConfigError("git {} failed: {}".format(arguments[0], result.stderr.strip() or result.stdout.strip()))
    return result


def value(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def load_rules() -> Dict[str, Any]:
    try:
        with RULES_PATH.open(encoding="utf-8") as stream:
            rules = yaml.safe_load(stream)
    except (OSError, yaml.YAMLError) as exc:
        raise SportConfigError("cannot read rule configuration: {}".format(exc))
    if not isinstance(rules, dict) or rules.get("schema_version") != 1:
        raise SportConfigError("unsupported rule configuration: {}".format(RULES_PATH))
    return rules


class WorkbookChecker:
    def __init__(self, workbook: Any, rules: Dict[str, Any]):
        self.workbook = workbook
        self.rules = rules
        self.findings: List[Finding] = []
        self._rows: Dict[str, List[Tuple[Any, ...]]] = {}

    def rows(self, sheet_name: str) -> List[Tuple[Any, ...]]:
        if sheet_name not in self._rows:
            self._rows[sheet_name] = list(self.workbook[sheet_name].iter_rows(values_only=True))
        return self._rows[sheet_name]

    def error(self, rule_id: str, message: str) -> None:
        self.findings.append(Finding("ERROR", rule_id, message))

    def warning(self, rule_id: str, message: str) -> None:
        self.findings.append(Finding("WARNING", rule_id, message))

    def header_index(self, sheet_name: str, column_name: str, header_row: Optional[int] = None) -> Optional[int]:
        rows = self.rows(sheet_name)
        candidates = [header_row] if header_row else [2, 1]
        candidates.extend(number for number in (2, 1) if number not in candidates)
        for number in candidates:
            if number is None or number < 1 or number > len(rows):
                continue
            for index, item in enumerate(rows[number - 1]):
                if value(item) == column_name:
                    return index
        return None

    def active_rows(self, sheet_name: str, spec: Dict[str, Any]) -> Iterable[Tuple[int, Tuple[Any, ...]]]:
        key = spec.get("key")
        if key == "row_label":
            return []
        key_index = self.header_index(sheet_name, key, spec.get("header_row"))
        if key_index is None:
            return []
        ignored = set()
        for item in spec.get("ignore_rows", []):
            column_index = self.header_index(sheet_name, item["column"], spec.get("header_row"))
            if column_index is not None:
                ignored.add((column_index, frozenset(value(entry) for entry in item.get("values", []))))
        start = int(spec.get("data_start_row", 3))
        range_spec = spec.get("row_range", {})
        end_marker = range_spec.get("end_marker")
        marker_column = int(range_spec.get("column", 1)) - 1
        records = []
        for row_number, row in enumerate(self.rows(sheet_name)[start - 1 :], start):
            key_value = value(row[key_index] if key_index < len(row) else None)
            if end_marker and value(row[marker_column] if marker_column < len(row) else None) == value(end_marker):
                break
            if not key_value:
                continue
            if any(value(row[index] if index < len(row) else None) in ignored_values for index, ignored_values in ignored):
                continue
            records.append((row_number, row))
        return records

    def check_required_sheets(self) -> None:
        for sheet_name in self.rules["scope"]["required_sheets"]:
            if sheet_name not in self.workbook.sheetnames:
                self.error("workbook.required_sheets", "missing Sheet '{}'".format(sheet_name))
        prefixes = tuple(self.rules["scope"].get("optional_sheet_prefixes", []))
        required = set(self.rules["scope"]["required_sheets"])
        for sheet_name in self.workbook.sheetnames:
            if sheet_name not in required and not sheet_name.startswith(prefixes):
                self.warning("workbook.unknown_sheet", "unknown Sheet '{}'".format(sheet_name))

    def check_row_markers(self) -> None:
        for layout in self.rules.get("sheet_layout", []):
            if layout.get("rule") != "row_markers" or layout["sheet"] not in self.workbook.sheetnames:
                continue
            column = int(layout.get("marker_column", 1)) - 1
            markers = [value(row[column] if column < len(row) else None) for row in self.rows(layout["sheet"])]
            start = layout["start_marker"]
            end = layout["end_marker"]
            if start not in markers or end not in markers:
                self.error(layout["id"], "{} marker missing in Sheet '{}'".format("start" if start not in markers else "end", layout["sheet"]))
            elif layout.get("start_before_end", False) and markers.index(start) >= markers.index(end):
                self.error(layout["id"], "start marker must precede end marker in Sheet '{}'".format(layout["sheet"]))

    def check_product_registry(self, spec: Dict[str, Any]) -> None:
        rows = self.rows("code_attr_product")
        range_spec = spec["row_range"]
        column = int(range_spec.get("column", 1)) - 1
        markers = [value(row[column] if column < len(row) else None) for row in rows]
        start = range_spec["start_marker"]
        end = range_spec["end_marker"]
        if start not in markers or end not in markers:
            return
        row_pattern = re.compile(spec["row_pattern"])
        for row_number, item in enumerate(markers[markers.index(start) + 1 : markers.index(end)], markers.index(start) + 2):
            if not item:
                self.error("code_attr_product.required_value", "Sheet code_attr_product row {} is blank".format(row_number))
            elif not row_pattern.fullmatch(item):
                self.error("code_attr_product.row_pattern", "Sheet code_attr_product row {} has invalid name '{}'".format(row_number, item))
            elif spec.get("require_registered_sheet_exists") and item not in self.workbook.sheetnames:
                self.error("code_attr_product.sheet_exists", "registered Sheet '{}' does not exist".format(item))

    def check_fields(self) -> None:
        for sheet_name, spec in self.rules.get("fields", {}).items():
            if sheet_name not in self.workbook.sheetnames:
                continue
            if sheet_name == "code_attr_product":
                self.check_product_registry(spec)
                continue
            rows = list(self.active_rows(sheet_name, spec))
            for column_name in spec.get("required_columns", []):
                if self.header_index(sheet_name, column_name, spec.get("header_row")) is None:
                    self.error("{}.required_columns".format(sheet_name), "Sheet {} missing column '{}'".format(sheet_name, column_name))
            if spec.get("key") == "row_label":
                labels = {value(row[0] if row else None) for row in self.rows(sheet_name)}
                for label in spec.get("row_labels", []):
                    if label not in labels:
                        self.error("{}.row_labels".format(sheet_name), "Sheet {} missing row label '{}'".format(sheet_name, label))
                continue
            for column_name, column_spec in spec.get("columns", {}).items():
                if column_name == "sport_columns":
                    continue
                index = self.header_index(sheet_name, column_name, spec.get("header_row"))
                if index is None:
                    continue
                seen: Dict[str, int] = {}
                allowed = column_spec.get("values")
                allowed_values = set(value(item) for item in allowed) if allowed is not None else None
                minimum = column_spec.get("constraints", {}).get("min")
                for row_number, row in rows:
                    item = row[index] if index < len(row) else None
                    normalized = value(item)
                    if column_spec.get("required_value") and not normalized:
                        self.error("{}.{}.required_value".format(sheet_name, column_name), "Sheet {} row {} column '{}' is blank".format(sheet_name, row_number, column_name))
                    if column_spec.get("unique") and normalized:
                        if normalized in seen:
                            self.error("{}.{}.unique".format(sheet_name, column_name), "Sheet {} rows {} and {} duplicate '{}'".format(sheet_name, seen[normalized], row_number, normalized))
                        else:
                            seen[normalized] = row_number
                    if allowed_values is not None and normalized not in allowed_values:
                        self.error("{}.{}.enum".format(sheet_name, column_name), "Sheet {} row {} column '{}' has unsupported value '{}'".format(sheet_name, row_number, column_name, normalized))
                    if minimum is not None and normalized:
                        try:
                            if int(item) < int(minimum):
                                self.error("{}.{}.min".format(sheet_name, column_name), "Sheet {} row {} column '{}' is below {}".format(sheet_name, row_number, column_name, minimum))
                        except (TypeError, ValueError):
                            self.error("{}.{}.integer".format(sheet_name, column_name), "Sheet {} row {} column '{}' must be an integer".format(sheet_name, row_number, column_name))
                    required_when = column_spec.get("required_when")
                    conditions = required_when if isinstance(required_when, list) else ([required_when] if required_when else [])
                    for condition in conditions:
                        if condition and self.matches_condition(sheet_name, row, spec, condition) and not normalized:
                            self.error("{}.{}.required_when".format(sheet_name, column_name), "Sheet {} row {} column '{}' is required by {}".format(sheet_name, row_number, column_name, condition["column"]))

    def matches_condition(self, sheet_name: str, row: Tuple[Any, ...], spec: Dict[str, Any], condition: Dict[str, Any]) -> bool:
        index = self.header_index(sheet_name, condition["column"], spec.get("header_row"))
        item = value(row[index] if index is not None and index < len(row) else None)
        if "equals" in condition:
            return item == value(condition["equals"])
        if "starts_with" in condition:
            return item.startswith(value(condition["starts_with"]))
        return bool(item) if condition.get("not_empty") else False

    def source_values(self, sheet_name: str, column_name: str) -> List[str]:
        spec = self.rules["fields"].get(sheet_name, {})
        index = self.header_index(sheet_name, column_name, spec.get("header_row"))
        if index is None:
            return []
        return [value(row[index] if index < len(row) else None) for _, row in self.active_rows(sheet_name, spec)]

    def check_relations(self) -> None:
        for relation in self.rules.get("relations", []):
            rule = relation["rule"]
            if rule == "foreign_key":
                source = self.source_values(relation["from"]["sheet"], relation["from"]["column"])
                target = set(self.source_values(relation["to"]["sheet"], relation["to"]["column"]))
                aliases = relation.get("aliases", {})
                normalize = relation.get("normalize")
                if normalize == "lower_trim":
                    target = {entry.lower() for entry in target}
                for item in source:
                    mapped = value(aliases.get(item, item))
                    compared = mapped.lower() if normalize == "lower_trim" else mapped
                    if compared and compared not in target:
                        self.error(relation["id"], "value '{}' is not present in {}.{}".format(item, relation["to"]["sheet"], relation["to"]["column"]))
            elif rule == "columns_match":
                source = self.source_values(relation["columns_from"]["sheet"], relation["columns_from"]["column"])
                target = {value(item) for item in self.rows(relation["sheet"])[int(relation.get("target_header_row", 1)) - 1] if value(item)}
                target -= set(relation.get("excluded_target_columns", []))
                aliases = relation.get("aliases", {})
                ignored = set(relation.get("ignored_source_values", []))
                for item in source:
                    if item and item not in ignored and value(aliases.get(item, item)) not in target:
                        self.error(relation["id"], "group column '{}' is missing from Sheet '{}'".format(item, relation["sheet"]))
            elif rule == "conditional_value":
                sheet_name = relation["sheet"]
                spec = self.rules["fields"][sheet_name]
                then_index = self.header_index(sheet_name, relation["then"]["column"], spec.get("header_row"))
                if then_index is None:
                    continue
                for row_number, row in self.active_rows(sheet_name, spec):
                    if self.matches_condition(sheet_name, row, spec, relation["when"]) and relation["then"].get("non_empty") and not value(row[then_index] if then_index < len(row) else None):
                        self.error(relation["id"], "Sheet {} row {} column '{}' is required".format(sheet_name, row_number, relation["then"]["column"]))

    def run(self) -> List[Finding]:
        self.check_required_sheets()
        self.check_row_markers()
        self.check_fields()
        self.check_relations()
        return self.findings


def check_xlsx(xlsx: Path) -> List[Finding]:
    try:
        # Some maintained sports.xlsx files have a stale worksheet dimension.
        # Normal mode reads the actual cells instead of trusting that dimension.
        workbook = load_workbook(xlsx, read_only=False, data_only=True)
    except Exception as exc:
        raise SportConfigError("cannot read sports.xlsx: {}".format(exc))
    try:
        return WorkbookChecker(workbook, load_rules()).run()
    finally:
        workbook.close()


def print_findings(findings: Sequence[Finding]) -> None:
    for finding in findings:
        print("{} [{}] {}".format(finding.severity, finding.rule_id, finding.message))


def run_check(repo: Path, xlsx: Path) -> CheckResult:
    if not xlsx_has_changes(repo, xlsx):
        print("SKIP: sports.xlsx unchanged")
        return CheckResult.SKIPPED
    findings = check_xlsx(xlsx)
    print_findings(findings)
    errors = [finding for finding in findings if finding.severity == "ERROR"]
    if errors:
        raise SportConfigError("sports.xlsx check failed: {} error(s)".format(len(errors)))
    print("PASS: sports.xlsx check passed")
    return CheckResult.PASSED


def run_generator(repo: Path, generator: Path) -> None:
    common_dir = generator.parent
    result = subprocess.run([sys.executable, generator.name], cwd=str(common_dir), check=False)
    if result.returncode:
        raise SportConfigError("sport_gen.py failed with exit {}".format(result.returncode))
    diff = git_command(git_root_for_path(common_dir), ["diff", "--stat"])
    print("PASS: sport_gen.py completed")
    print(diff.stdout.rstrip() or "git diff --stat: no unstaged diff")


def run(arguments: Sequence[str], cwd: Optional[Path] = None) -> int:
    parser = build_parser()
    args = parser.parse_args(normalize_action_aliases(arguments))
    start = Path(args.repo) if args.repo else (cwd or Path.cwd())
    try:
        repo = find_repo(start)
        xlsx, generator = require_layout(repo)
        result = run_check(repo, xlsx)
        if args.command == "gen":
            if result is CheckResult.SKIPPED:
                print("SKIP: generation not run because sports.xlsx is unchanged")
            else:
                run_generator(repo, generator)
        return 0
    except SportConfigError as exc:
        parser.error(str(exc))
    return 2


def main() -> int:
    return run(sys.argv[1:])


if __name__ == "__main__":
    raise SystemExit(main())
