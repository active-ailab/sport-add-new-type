# -*- coding: utf-8 -*-
"""
变更计划数据结构。

设计约定：
- 一个变更操作（如新增运动）产出单个 ChangePlan。
- 计划内含三类原子项：
    * CellSet   —— 覆盖/清空单元格。坐标为"最终坐标"（含同一计划内行/列插入的偏移）。
    * RowInsert / RowDelete —— 行级插入/删除。at 为插入/删除起点行号（1 基）。
    * ColInsert / ColDelete —— 列级插入/删除。at 为插入/删除起点列号（1 基）。
- apply(wb) 顺序：先处理所有行/列 增删（跨表顺序任意、同表按 at 倒序避免坐标漂移），
  最后统一执行 CellSet（坐标为最终坐标）。
- describe() 生成人类可读的逐项 diff，用于 dry-run 预览与 UI 展示。
"""
from __future__ import annotations

from dataclasses import dataclass, field
from typing import Optional

import openpyxl
from openpyxl.utils import get_column_letter


@dataclass
class CellSet:
    sheet: str
    row: int
    col: int
    new: object
    old: object = None          # 预览时由 ops 读取原值填入
    reason: str = ""

    def describe(self):
        old_s = "" if self.old is None else str(self.old)
        new_s = "" if self.new is None else str(self.new)
        mark = "=" if old_s == new_s else "→"
        return f"  {self.sheet}!{get_column_letter(self.col)}{self.row}  {old_s!r} {mark} {new_s!r}  ({self.reason})"


@dataclass
class RowInsert:
    sheet: str
    at: int                      # 在该行号之前插入
    count: int = 1
    reason: str = ""

    def describe(self):
        return f"  {self.sheet}: 在第 {self.at} 行前插入 {self.count} 行  ({self.reason})"


@dataclass
class RowDelete:
    sheet: str
    at: int
    count: int = 1
    reason: str = ""

    def describe(self):
        return f"  {self.sheet}: 删除第 {self.at}..{self.at + self.count - 1} 行  ({self.reason})"


@dataclass
class ColInsert:
    sheet: str
    at: int
    count: int = 1
    reason: str = ""

    def describe(self):
        return f"  {self.sheet}: 在第 {get_column_letter(self.at)} 列前插入 {self.count} 列  ({self.reason})"


@dataclass
class ColDelete:
    sheet: str
    at: int
    count: int = 1
    reason: str = ""

    def describe(self):
        return f"  {self.sheet}: 删除第 {get_column_letter(self.at)}..{get_column_letter(self.at + self.count - 1)} 列  ({self.reason})"


@dataclass
class ChangePlan:
    title: str
    items: list = field(default_factory=list)

    def add(self, item):
        self.items.append(item)
        return self

    @property
    def empty(self):
        return not self.items

    def describe(self) -> str:
        lines = [f"=== {self.title} ==="]
        if not self.items:
            lines.append("  (无变更)")
            return "\n".join(lines)
        for it in self.items:
            lines.append(it.describe())
        return "\n".join(lines)

    def apply(self, wb: openpyxl.Workbook):
        """按『先行列增删、后单元格覆盖』执行。行列增删按 同表 at 倒序，避免坐标漂移。"""
        # 行/列增删：先倒序处理同表内的同向操作
        by_sheet = {}
        for it in self.items:
            if isinstance(it, (RowInsert, RowDelete, ColInsert, ColDelete)):
                by_sheet.setdefault(it.sheet, []).append(it)
        for sheet, ops in by_sheet.items():
            ws = wb[sheet]
            # 倒序执行（后插入点先执行）
            for it in sorted(ops, key=lambda x: x.at, reverse=True):
                if isinstance(it, RowInsert):
                    ws.insert_rows(it.at, it.count)
                elif isinstance(it, RowDelete):
                    ws.delete_rows(it.at, it.count)
                elif isinstance(it, ColInsert):
                    ws.insert_cols(it.at, it.count)
                elif isinstance(it, ColDelete):
                    ws.delete_cols(it.at, it.count)
        # 单元格覆盖（坐标为最终坐标）
        for it in self.items:
            if isinstance(it, CellSet):
                wb[it.sheet].cell(row=it.row, column=it.col).value = it.new
        return self
